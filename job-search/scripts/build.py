"""Filter, score against Yashik's CV, and write the Excel workbook."""
import json, re, sys, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

RAW, INDEED, OUT = sys.argv[1:4]
jobs = json.load(open(RAW))
for line in open(INDEED, encoding="utf-8"):
    t, c, l, p, et, u = line.rstrip("\n").split("\t")
    jobs.append({"source": "Indeed", "id": "in-" + u[-12:], "title": t, "company": c, "location": l,
                 "posted": p, "employment_type": et, "url": u, "description": ""})

EXCL_TITLE = re.compile(r"\b(senior|sr\.?|lead|leiter|leitung|principal|head|staff|director|direktor|manager|"
                        r"architect|architekt|werkstudent\w*|working student|student\w*|intern|internship|praktik\w*|"
                        r"thesis|abschlussarbeit|masterarbeit|bachelorarbeit|ausbildung|azubi|auszubildende\w*|dual\w*|"
                        r"minijob|teilzeit|part[- ]time|freelance\w*|freiberuf\w*|expert|experienced|chief|vp|"
                        r"teamlead|tech lead|cto|professor|postdoc|phd|doktorand\w*|mid[- ]level|"
                        r"vertrieb\w*|sales|account|kaufm\w*|recruit\w*|elektriker|mechatron\w*|techniker|"
                        r"monteur|servicetechniker|lager\w*|fahrer|pflege\w*)\b", re.I)
TECH_TITLE = re.compile(r"(develop|entwickl|engineer|ingenieur|software|programm|\bdata\b|daten|\bai\b|\bki\b|"
                        r"machine learning|\bml\b|\bllm|genai|devops|cloud|front.?end|back.?end|full.?stack|"
                        r"\bweb|python|java|typescript|react|angular|node|blockchain|solidity|\bios\b|swift|"
                        r"flutter|mobile|\bapp\b|platform|\bsre\b|security|it[- ]consultant|informatik|"
                        r"scientist|analytics|automation|automatisierung|rag\b|\bqa\b|test)", re.I)
JUNIOR = re.compile(r"(junior|\bjr\.?\b|entry|einstieg|einsteiger|berufseinsteiger|graduate|absolvent|"
                    r"trainee|young professional|new grad|associate)", re.I)
NONTECH_TITLE = re.compile(r"(elektro|maschinenbau|verfahrens|bau\w*ingenieur|konstrukt|vertrieb|sales|"
                           r"kalkulat|hardware|hochfrequenz|fpga|pcb|chemie|mechani|hvac|tga|"
                           r"projektingenieur|prüfingenieur|qualitätsingenieur|bauleit|statik|"
                           r"versorgungstechnik|fertigung|produktion|process engineer|netzwerk|network engineer|"
                           r"system\s?ingenieur|systemingenieur|automatisierungsingenieur|automation engineer \(plc|\bplc\b|\bsps\b|"
                           r"pharma|\bhmi\b|kommunikationstechnik|nachrichtentechnik|it[- ]operations|administrator|support|"
                           r"helpdesk|service desk|sap (?:basis|hcm)|abap|salesforce admin)", re.I)

YEARS = re.compile(r"(?:(?:mind(?:estens|\.)?|min(?:imum|\.)?|at least|über|over)\s*)?(\d{1,2})\s*\+?\s*"
                   r"(?:[-–]\s*\d{1,2}\s*)?(?:\+\s*)?(?:years?|yrs|jahre?n?)\b(?![^.\n]{0,15}(?:alt|old|garantie|warranty))", re.I)

SKILLS = {  # CV skill -> (regex, weight)
    "Python": (r"\bpython\b", 3), "TypeScript": (r"\btypescript\b|\bts\b", 3), "JavaScript": (r"javascript|\bjs\b", 2),
    "React": (r"\breact(?:\.js|js)?\b", 3), "Angular": (r"\bangular\b", 3), "Node.js": (r"\bnode(?:\.js|js)?\b", 3),
    "LLMs/GenAI": (r"\bllms?\b|large language model|generative ai|gen ?ai|\bgpt\b|openai|anthropic|claude", 4),
    "RAG": (r"\brag\b|retrieval[- ]augmented", 4), "LangChain/LangGraph": (r"langchain|langgraph", 4),
    "LlamaIndex": (r"llama.?index", 4), "AI Agents": (r"\bagent(?:s|ic)\b|ki-agent", 3),
    "Vector DBs": (r"vector (?:db|database|store)|chroma|qdrant|pinecone|weaviate|pgvector", 3),
    "Search (BM25/Meilisearch/Elastic)": (r"meilisearch|elasticsearch|opensearch|bm25|\bsolr\b", 2),
    "Machine Learning": (r"machine learning|\bml\b|maschinelles lernen|pytorch|tensorflow|scikit|hugging ?face", 2),
    "Docker": (r"\bdocker\b|container", 2), "Kubernetes": (r"kubernetes|\bk8s\b", 1),
    "PostgreSQL": (r"postgres", 2), "SQL": (r"\bsql\b", 1), "GraphQL": (r"graphql", 1),
    "Fastify/Express": (r"fastify|express(?:\.js|js)?\b|nestjs|nest\.js", 2), "REST APIs": (r"\brest\b|restful|\bapis?\b", 1),
    "CI/CD / GitHub Actions": (r"ci/cd|github actions|gitlab ci|jenkins|continuous integration", 1),
    "AWS/Azure/Cloud": (r"\baws\b|azure|\bgcp\b|google cloud|oracle cloud|\bcloud\b", 1), "Linux": (r"\blinux\b", 1),
    "Data pipelines / ETL": (r"data pipeline|etl\b|elt\b|datenpipeline|scraping|crawler|data integration|node-red", 2),
    "Blockchain/Solidity/Web3": (r"blockchain|solidity|web3|ethereum|smart contract|\bdefi\b", 3),
    "Swift/iOS": (r"\bswift\b|swiftui|\bios\b", 2), "Flutter/Dart": (r"flutter|\bdart\b", 2),
    "Redux": (r"\bredux\b", 1), "Testing (pytest/Vitest)": (r"pytest|vitest|jest|unit test", 1),
    "Security / SAST": (r"\bsast\b|sonarqube|appsec|application security|secure coding", 2),
    "Full-stack": (r"full.?stack", 2), "Monorepo/pnpm": (r"monorepo|pnpm|turborepo|nx\b", 1),
}
SKILLS = {k: (re.compile(r, re.I), w) for k, (r, w) in SKILLS.items()}
MAXW = 22.0

CATS = [("AI / ML / LLM", r"\bai\b|\bki\b|machine learning|\bml\b|\bllm|genai|generative|data scien|nlp|computer vision|deep learning|rag\b|agent"),
        ("Blockchain / Web3", r"blockchain|solidity|web3|crypto|smart contract"),
        ("Data Engineering / Analytics", r"\bdata\b|daten|analytics|\bbi\b|etl"),
        ("DevOps / Cloud / Platform", r"devops|cloud|platform|\bsre\b|infrastru|kubernetes|site reliab"),
        ("Security", r"security|sicherheit|cyber|appsec"),
        ("Mobile (iOS / Flutter)", r"\bios\b|android|mobile|flutter|swift|\bapp\b.*entwickl|app develop"),
        ("Full-Stack", r"full.?stack|fullstack"), ("Frontend", r"front.?end|\bui\b|react|angular|vue|web.?entwick|web develop"),
        ("Backend", r"back.?end|python|java\b|node|\.net|c#|golang|\bgo\b|php|api"),
        ("QA / Test Automation", r"\bqa\b|test|quality assurance"),
        ("Software Engineering (general)", r".")]
CATS = [(n, re.compile(r, re.I)) for n, r in CATS]

GER_REQ = re.compile(r"(sehr gute[n]? deutsch|fließend\w* deutsch|deutsch\w* fließend|verhandlungssicher\w*|"
                     r"deutsch\w* (?:in wort und schrift|auf c1|c1|c2|muttersprach)|(?:fluent|excellent|very good|strong|"
                     r"business[- ]fluent|native)\s+(?:\w+\s+)?german|german\s*\(?(?:c1|c2|native|fluent)|german (?:is )?(?:a )?(?:must|required|mandatory)|"
                     r"proficiency in german|german and english (?:are|is) (?:required|essential|a must)|deutsch und englisch|"
                     r"(?:gute|sehr gut\w*) deutsch- und englischkenntnisse|deutschkenntnisse)", re.I)
GER_PLUS = re.compile(r"german (?:is |would be )?(?:a )?(?:plus|bonus|nice|advantage)|german (?:language )?skills (?:are|is) (?:a )?(?:plus|bonus|advantage)|"
                      r"nice to have[^.\n]{0,40}german|german[^.\n]{0,40}nice to have|(?:basic|some) german", re.I)
NO_GER = re.compile(r"no german (?:required|needed|necessary)|german (?:is )?not (?:required|necessary|needed)|english[- ]speaking (?:team|environment|company)|"
                    r"(?:company|working|office) language is english|english is our (?:company|working) language|"
                    r"without german", re.I)
VISA = re.compile(r"visa|relocation|blue card|sponsorship|umzugsunterstützung|relocat", re.I)
REMOTE = re.compile(r"\bremote\b|home.?office|mobiles arbeiten|mobile work|work from home|homeoffice|ortsunabhängig", re.I)
HYBRID = re.compile(r"hybrid|teilweise remote|flexible[s]? arbeiten|days? (?:in the )?office", re.I)
DE_WORDS = re.compile(r"\b(und|die|der|wir|für|mit|eine?n?|sie|bei|uns|deine?|ihre?|sind|auf)\b", re.I)
EN_WORDS = re.compile(r"\b(and|the|we|for|with|you|our|your|are|of|to|is)\b", re.I)

BIG = ["sap", "siemens", "bosch", "mercedes", "bmw", "volkswagen", "vw ", "audi", "porsche", "allianz", "deutsche telekom", "telekom",
       "t-systems", "accenture", "capgemini", "ibm", "microsoft", "amazon", "aws", "google", "meta", "apple", "zalando", "delivery hero",
       "check24", "deloitte", "pwc", "kpmg", "ey ", "ernst & young", "infineon", "continental", "zf ", "basf", "bayer", "deutsche bank",
       "commerzbank", "dhl", "deutsche post", "lufthansa", "db systel", "deutsche bahn", "otto", "rewe", "lidl", "schwarz", "kaufland",
       "aldi", "trumpf", "bertelsmann", "arvato", "airbus", "thales", "ericsson", "nokia", "vodafone", "telefonica", "o2", "1&1", "united internet",
       "ionos", "adesso", "msg ", "materna", "atos", "sopra steria", "cgi", "nttdata", "ntt data", "tcs", "infosys", "wipro", "hcl", "cognizant",
       "freenet", "kion", "heidelberg", "henkel", "merck", "siemens healthineers", "rohde", "zeiss", "festo", "sick ag", "beckhoff", "wago",
       "miele", "dspace", "hella", "forvia", "diebold", "wincor", "phoenix contact", "bertrandt", "akkodis", "alten", "ferchau", "valtech",
       "booking", "jetbrains", "celonis", "personio", "n26", "trade republic", "hellofresh", "auto1", "scout24", "idealo", "about you",
       "flix", "tier", "gorillas", "wolt", "uber", "spotify", "ebay", "paypal", "mckinsey", "bcg", "boston consulting", "porsche digital",
       "mercedes-benz", "e.on", "rwe", "enbw", "vattenfall", "lhind", "munich re", "ergo", "hdi", "talanx", "generali", "axa", "dkb", "ing ",
       "n-ergie", "datev", "dm-drogerie", "mediamarkt", "ceconomy", "aldi süd", "eon", "statista", "hubspot", "adobe", "oracle", "salesforce",
       "cisco", "intel", "nvidia", "qualcomm", "amd", "huawei", "bmw group", "continental ag", "schaeffler", "mahle", "thyssenkrupp", "deutsche börse",
       "dzbank", "dz bank", "union investment", "ceva", "emerson", "honeywell", "abb", "schneider electric", "philips", "nxp", "bosch rexroth"]


def text(j):
    return f"{j.get('title','')}\n{j.get('description','')}"


def min_years(d):
    ys = []
    for m in YEARS.finditer(d or ""):
        n = int(m.group(1))
        ctx = d[max(0, m.start() - 60): m.end() + 60].lower()
        if 1 <= n <= 15 and any(w in ctx for w in ("experience", "erfahrung", "berufserfahrung", "work")):
            ys.append(n)
    return min(ys) if ys else None


def lang(d):
    de, en = len(DE_WORDS.findall(d)), len(EN_WORDS.findall(d))
    if de + en < 15:
        return ""
    return "German" if de > en else "English"


def german_req(j, lg):
    d = j.get("description") or ""
    if NO_GER.search(d):
        return "Not required"
    if GER_REQ.search(d):
        return "Required (fluent)"
    if GER_PLUS.search(d):
        return "Plus / nice-to-have"
    if lg == "German":
        return "Likely required (ad in German)"
    if lg == "English":
        return "Not mentioned (ad in English)"
    t = j.get("title", "")
    if re.search(r"\(m/w/d\)|entwickler|\bw/m/d\b", t, re.I) and not re.search(r"engineer|developer", t, re.I):
        return "Likely required (German title)"
    return "Unknown"


def keep(j):
    t = j.get("title", "")
    if not t or EXCL_TITLE.search(t) or not TECH_TITLE.search(t) or NONTECH_TITLE.search(t):
        return False
    et = (j.get("employment_type") or "").lower()
    if any(x in et for x in ("part", "intern", "contract", "temporary", "volunteer", "other", "freelance")) and "full" not in et:
        return False
    if j["source"] == "LinkedIn":
        s = (j.get("seniority") or "").lower()
        if s and s not in ("entry level", "associate", "not applicable", ""):
            return False
        if "Germany" not in j.get("location", "") and "Deutschland" not in j.get("location", ""):
            return False
    if j["source"] == "Arbeitnow":
        d = j.get("description", "")
        if re.search(r"experienced|\bmid\b|berufserfahren|senior", et) and not JUNIOR.search(t):
            return False
        loc = j.get("location", "")
        if re.search(r"\b(usa|united states|uk|london|austria|österreich|schweiz|switzerland|wien|zürich)\b", loc, re.I):
            return False
        if not (JUNIOR.search(t) or re.search(r"berufseinsteiger|erste berufserfahrung|first (?:professional|working )?experience|"
                                              r"(?:0|1)\s*[-–]\s*[23]\s*(?:years|jahre)|graduates?\b|absolvent|career start|berufsstart|"
                                              r"junior", d, re.I)):
            return False
    my = min_years(j.get("description", ""))
    if my is not None and my >= 4 and not JUNIOR.search(t):
        return False
    return True


def score(j):
    d = text(j)
    hits = [k for k, (r, w) in SKILLS.items() if r.search(d)]
    raw = sum(SKILLS[k][1] for k in hits)
    s = min(raw / MAXW, 1.0) * 65
    t = j["title"]
    if re.search(r"\bai\b|\bki\b|llm|genai|machine learning|\bml\b|agent|rag", t, re.I): s += 22
    elif re.search(r"full.?stack|python|typescript|node|react|angular|backend|back-end", t, re.I): s += 16
    elif re.search(r"software|develop|entwickl|engineer", t, re.I): s += 7
    if JUNIOR.search(t): s += 8
    if not j.get("description"):  # no description available -> estimate from title only
        s = 62 if re.search(r"\bai\b|\bki\b|llm|genai|machine learning|\bml\b|agent", t, re.I) else \
            58 if re.search(r"full.?stack|python|typescript|node|react|angular|backend|frontend", t, re.I) else \
            48 if re.search(r"software|develop|entwickl|engineer|data|cloud", t, re.I) else 40
        s += 5 if JUNIOR.search(t) else 0
    return round(min(s, 100)), hits


def mode(j):
    d = text(j) + " " + j.get("location", "")
    if j.get("remote_flag") is True or re.search(r"\bremote\b", j.get("location", ""), re.I): return "Remote"
    if HYBRID.search(d): return "Hybrid"
    if REMOTE.search(d): return "Remote / Hybrid possible"
    return "On-site / not stated"


def company_size(c):
    lc = f" {c.lower()} "
    return "Large enterprise / MNC" if any((" " + b) in lc or lc.strip().startswith(b.strip()) for b in BIG) else "SME / Startup (unverified)"


def norm_key(j):
    return (re.sub(r"\W+", "", j["title"].lower())[:60], re.sub(r"\W+", "", j["company"].lower())[:30])


rows, seen = [], set()
order = {"LinkedIn": 0, "Arbeitnow": 1, "Bundesagentur für Arbeit": 2, "Indeed": 3}
for j in sorted(jobs, key=lambda x: order.get(x["source"], 9)):
    if not keep(j):
        continue
    k = norm_key(j)
    if k in seen:
        continue
    seen.add(k)
    sc, hits = score(j)
    if sc < 20:
        continue
    d = j.get("description", "")
    lg = lang(d)
    cat = next(n for n, r in CATS if r.search(j["title"]))
    city = (j.get("location") or "").split(",")[0].strip()
    rows.append({
        "Fit Score": sc, "Job Title": j["title"], "Company": j["company"], "Company Type": company_size(j["company"]),
        "City": city, "Location (full)": j.get("location", ""), "Work Mode": mode(j), "Category": cat,
        "Seniority": j.get("seniority") or ("Junior" if JUNIOR.search(j["title"]) else "Entry/Junior (filtered)"),
        "Employment Type": j.get("employment_type") or "Full-time", "Min. Years Exp (detected)": min_years(d),
        "German Requirement": german_req(j, lg), "Ad Language": lg or "n/a",
        "Visa / Relocation Mentioned": "Yes" if VISA.search(d) else "", "Salary": j.get("salary", ""),
        "Matched CV Skills": ", ".join(hits),
        "Score Basis": "Full description" if d else "Title only", "# Skills Matched": len(hits),
        "Industry": j.get("industry", "") or j.get("tags", ""), "Job Function": j.get("job_function", ""),
        "Applicants": j.get("applicants", ""), "Posted": (j.get("posted") or "")[:10], "Source": j["source"],
        "Search Query": j.get("query", ""), "Job URL": j["url"], "Company Page": j.get("company_url", ""),
        "Description (excerpt)": re.sub(r"\s+", " ", d)[:600],
        "Status": "Not applied", "Applied On": "", "Notes": "",
    })

rows.sort(key=lambda r: (-r["Fit Score"], r["Posted"] and -int(r["Posted"].replace("-", "") or 0)))
for i, r in enumerate(rows, 1):
    r["Tier"] = "A - Strong match" if r["Fit Score"] >= 70 else "B - Good match" if r["Fit Score"] >= 50 else "C - Possible"

COLS = ["#", "Tier", "Fit Score", "Job Title", "Company", "Company Type", "City", "Location (full)", "Work Mode", "Category",
        "Seniority", "Employment Type", "Min. Years Exp (detected)", "German Requirement", "Ad Language",
        "Visa / Relocation Mentioned", "Salary", "# Skills Matched", "Matched CV Skills", "Score Basis", "Industry", "Job Function",
        "Applicants", "Posted", "Source", "Search Query", "Job URL", "Company Page", "Description (excerpt)",
        "Status", "Applied On", "Notes"]
WIDTH = {"#": 6, "Tier": 16, "Fit Score": 9, "Job Title": 48, "Company": 28, "Company Type": 22, "City": 16, "Location (full)": 28,
         "Work Mode": 20, "Category": 24, "Seniority": 16, "Employment Type": 14, "Min. Years Exp (detected)": 11,
         "German Requirement": 26, "Ad Language": 11, "Visa / Relocation Mentioned": 11, "Salary": 18, "# Skills Matched": 9,
         "Matched CV Skills": 50, "Score Basis": 14, "Industry": 26, "Job Function": 22, "Applicants": 18, "Posted": 11, "Source": 14,
         "Search Query": 24, "Job URL": 18, "Company Page": 14, "Description (excerpt)": 70, "Status": 13, "Applied On": 11, "Notes": 30}

HDR = PatternFill("solid", fgColor="1A5276")
wb = Workbook()


def sheet(ws, data, name):
    ws.append(COLS)
    for i, r in enumerate(data, 1):
        ws.append([i if c == "#" else r.get(c) for c in COLS])
    ui, ci = COLS.index("Job URL") + 1, COLS.index("Company Page") + 1
    for row in ws.iter_rows(min_row=2):
        for idx in (ui, ci):
            cell = row[idx - 1]
            if cell.value:
                cell.hyperlink = cell.value
                cell.value = "Open posting" if idx == ui else "Company"
                cell.style = "Hyperlink"
        row[COLS.index("Description (excerpt)")].alignment = Alignment(wrap_text=False)
    for i, c in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTH[c]
        h = ws.cell(1, i); h.font = Font(bold=True, color="FFFFFF"); h.fill = HDR
        h.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "E2"
    ref = f"A1:{get_column_letter(len(COLS))}{len(data) + 1}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    sc = get_column_letter(COLS.index("Fit Score") + 1)
    ws.conditional_formatting.add(f"{sc}2:{sc}{len(data) + 1}",
                                  ColorScaleRule(start_type="num", start_value=20, start_color="F8696B",
                                                 mid_type="num", mid_value=55, mid_color="FFEB84",
                                                 end_type="num", end_value=90, end_color="63BE7B"))


ws = wb.active; ws.title = "All Jobs"; sheet(ws, rows, "AllJobs")
top = [r for r in rows if r["Fit Score"] >= 60]
sheet(wb.create_sheet("Top Matches (60+)"), top, "TopMatches")
eng = [r for r in rows if r["German Requirement"] in ("Not required", "Not mentioned (ad in English)", "Plus / nice-to-have")]
sheet(wb.create_sheet("English-Friendly"), eng, "EnglishFriendly")
ai = [r for r in rows if r["Category"] == "AI / ML / LLM"]
sheet(wb.create_sheet("AI-ML-LLM Roles"), ai, "AIRoles")
pb = [r for r in rows if re.search(r"paderborn|bielefeld|gütersloh|guetersloh|dortmund|münster|muenster|kassel|lippstadt|"
                                  r"detmold|bad salzuflen|hamm|soest|herford|minden|lemgo|osnabrück|hannover|köln|cologne|düsseldorf|"
                                  r"duesseldorf|essen|bochum|duisburg|wuppertal|nrw|north rhine", r["Location (full)"], re.I)]
sheet(wb.create_sheet("Near Paderborn - NRW"), pb, "NearPaderborn")

s = wb.create_sheet("Summary", 0)
s["A1"] = "Entry / Junior Full-Time Tech Jobs in Germany — matched to Yashikkumar Khunt's CV"; s["A1"].font = Font(bold=True, size=14, color="1A5276")
s["A2"] = f"Generated {datetime.date.today().isoformat()} · {len(rows)} unique jobs after filtering {len(jobs)} raw postings"
r0 = 4


def block(title, counter, r):
    s.cell(r, 1, title).font = Font(bold=True, color="FFFFFF"); s.cell(r, 1).fill = HDR
    s.cell(r, 2, "Jobs").font = Font(bold=True, color="FFFFFF"); s.cell(r, 2).fill = HDR
    for k, v in counter:
        r += 1; s.cell(r, 1, k); s.cell(r, 2, v)
    return r + 2


r0 = block("Tier", sorted(collections.Counter(r["Tier"] for r in rows).items()), r0)
r0 = block("Category", collections.Counter(r["Category"] for r in rows).most_common(), r0)
r0 = block("German requirement", collections.Counter(r["German Requirement"] for r in rows).most_common(), r0)
r0 = block("Work mode", collections.Counter(r["Work Mode"] for r in rows).most_common(), r0)
r0 = block("Company type", collections.Counter(r["Company Type"] for r in rows).most_common(), r0)
r0 = block("Top 25 cities", collections.Counter(r["City"] for r in rows).most_common(25), r0)
r0 = block("Source", collections.Counter(r["Source"] for r in rows).most_common(), r0)
r0 = block("Most-hiring companies (top 25)", collections.Counter(r["Company"] for r in rows).most_common(25), r0)
notes = [
    "HOW TO READ",
    "Fit Score (0-100): weighted overlap between the job title/description and the skills in the CV (Python, TypeScript, React/Angular, Node, LLMs, RAG, LangChain/LangGraph/LlamaIndex, Docker, PostgreSQL, etc.) plus a bonus for AI/full-stack titles and explicit junior wording.",
    "Tiers: A = 70+, B = 50-69, C = below 50. Indeed rows have no description, so they're scored on the title only.",
    "German Requirement is detected from the ad text (e.g. 'fließend Deutsch', 'German C1'). Ads written in German usually expect strong German; the CV lists German A2.",
    "Company Type is a heuristic: a known-name list marks large enterprises/MNCs; everything else is SME/startup (unverified).",
    "Filters applied: full-time only; Germany only; excluded senior/lead/principal/manager/architect, Werkstudent, internships, theses, Ausbildung, part-time, freelance and non-software engineering titles; excluded ads that ask for 4+ years unless the title says junior.",
    "Sources: LinkedIn public job search (Entry level + Associate, full-time, posted in the last 30 days), Arbeitnow job board API, Indeed (via connector).",
    "Status / Applied On / Notes columns are left blank for tracking applications.",
]
for i, n in enumerate(notes):
    c = s.cell(4 + i, 4, n); c.alignment = Alignment(wrap_text=True, vertical="top")
    if i == 0: c.font = Font(bold=True)
s.column_dimensions["A"].width = 42; s.column_dimensions["B"].width = 8; s.column_dimensions["D"].width = 110
wb.save(OUT)
json.dump(rows, open(OUT.replace(".xlsx", ".json"), "w"), ensure_ascii=False)
print("rows", len(rows), "top", len(top), "eng", len(eng), "ai", len(ai), "nrw", len(pb))
print(collections.Counter(r["Source"] for r in rows))
